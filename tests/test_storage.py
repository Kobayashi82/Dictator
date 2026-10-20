import json

import pytest

from dictado.doc import storage
from dictado.doc.model import Paragraph


def sample():
    return [
        Paragraph(raw="El viejo y el mar", status="closed", style="title"),
        Paragraph(raw="era un viejo que pescaba solo", corrected="Era un viejo que pescaba solo.", status="corrected"),
        Paragraph(raw="Verso uno\nverso dos", status="rejected", mode="poesia", reason="cambia el número de versos"),
        Paragraph(raw="texto original", corrected="Texto a mano.", status="edited", user_edited=True),
        Paragraph(raw="sigo hablando", status="dictating"),
    ]


def test_roundtrip_keeps_every_field(tmp_path):
    path = tmp_path / "doc.dictado"
    paragraphs = sample()
    storage.save(path, paragraphs, "poesia", ["Santiago", "Manolín"])
    loaded, mode, glossary = storage.load(path)
    assert mode == "poesia"
    assert glossary == ["Santiago", "Manolín"]
    assert [p.to_dict() for p in loaded] == [p.to_dict() for p in paragraphs]


def test_correcting_becomes_pending_on_load(tmp_path):
    path = tmp_path / "doc.dictado"
    storage.save(path, [Paragraph(raw="a", status="correcting")], "prosa", [])
    loaded, _, _ = storage.load(path)
    assert loaded[0].status == "pending"


def test_created_is_preserved_on_resave(tmp_path):
    path = tmp_path / "doc.dictado"
    storage.save(path, [], "prosa", [])
    created = json.loads(path.read_text(encoding="utf-8"))["created"]
    storage.save(path, sample(), "prosa", [])
    assert json.loads(path.read_text(encoding="utf-8"))["created"] == created


@pytest.mark.parametrize("content", ["no es json", json.dumps({"format": "otro"}), json.dumps([1, 2])])
def test_invalid_files_raise(tmp_path, content):
    path = tmp_path / "x.dictado"
    path.write_text(content, encoding="utf-8")
    with pytest.raises(storage.StorageError):
        storage.load(path)


def test_newer_version_is_rejected(tmp_path):
    path = tmp_path / "x.dictado"
    path.write_text(json.dumps({"format": "dictado", "version": 99, "paragraphs": []}), encoding="utf-8")
    with pytest.raises(storage.StorageError):
        storage.load(path)


def test_visible_text():
    p = sample()
    assert p[1].text == "Era un viejo que pescaba solo."
    p[1].show_raw = True
    assert p[1].text == "era un viejo que pescaba solo"
    assert p[2].text == "Verso uno\nverso dos"  # rechazado: se ve el crudo
    assert p[3].text == "Texto a mano."


def test_export_text_and_markdown():
    p = sample()
    txt = storage.export_text(p)
    assert txt.split("\n\n")[1] == "Era un viejo que pescaba solo."
    md = storage.export_markdown(p)
    assert md.startswith("# El viejo y el mar\n\n")
    assert "Verso uno  \nverso dos" in md


def test_export_docx(tmp_path):
    from docx import Document

    path = tmp_path / "out.docx"
    notes = Paragraph(raw="- pan\n- leche", status="closed", mode="notas")
    storage.export(path, sample() + [notes])
    doc = Document(str(path))
    texts = [par.text for par in doc.paragraphs]
    assert texts[0] == "El viejo y el mar"
    assert doc.paragraphs[0].style.name.startswith("Heading")
    assert "Era un viejo que pescaba solo." in texts
    assert "pan" in texts and "leche" in texts


def test_import_text_and_markdown(tmp_path):
    md = tmp_path / "a.md"
    md.write_text("# Título\n\nPrimer párrafo\nsigue.\n\n\nSegundo  \ncon salto.\n", encoding="utf-8")
    ps = storage.import_text(md)
    assert [p.style for p in ps] == ["title", "normal", "normal"]
    assert ps[0].text == "Título"
    assert ps[2].text == "Segundo\ncon salto."
    assert all(p.user_edited and p.status == "edited" for p in ps)

    txt = tmp_path / "b.txt"
    txt.write_bytes("Uno.\r\n\r\nDós.\r\n".encode("cp1252"))
    assert [p.text for p in storage.import_text(txt)] == ["Uno.", "Dós."]
