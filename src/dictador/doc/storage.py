"""Guardar y abrir documentos `.dictador` (y los antiguos `.dictado`), exportar e importar."""

from __future__ import annotations

import json
import os
import re
from datetime import datetime
from pathlib import Path

from dictador.doc.model import Paragraph

FORMAT = "dictador"
OLD_FORMATS = ("dictado",)  # versiones anteriores del programa, que se llamaba «Dictado»
VERSION = 1
EXTENSION = ".dictador"


class StorageError(Exception):
    pass


def _atomic_write(path: Path, data: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    tmp.write_text(data, encoding="utf-8")
    os.replace(tmp, path)


# -- formato nativo ---------------------------------------------------------------------------------------------------


def to_json(paragraphs: list[Paragraph], mode: str, glossary: list[str], created: str | None = None) -> str:
    now = datetime.now().isoformat(timespec="seconds")
    data = {
        "format": FORMAT,
        "version": VERSION,
        "created": created or now,
        "modified": now,
        "mode": mode,
        "glossary": list(glossary),
        "paragraphs": [p.to_dict() for p in paragraphs],
    }
    return json.dumps(data, ensure_ascii=False, indent=2)


def save(path: str | Path, paragraphs: list[Paragraph], mode: str, glossary: list[str]) -> None:
    path = Path(path)
    created = None
    if path.exists():
        try:
            created = json.loads(path.read_text(encoding="utf-8")).get("created")
        except (OSError, ValueError, AttributeError):
            pass
    _atomic_write(path, to_json(paragraphs, mode, glossary, created))


def load(path: str | Path) -> tuple[list[Paragraph], str, list[str]]:
    """Devuelve (párrafos, modo, glosario del documento)."""
    try:
        data = json.loads(Path(path).read_text(encoding="utf-8"))
    except OSError as e:
        raise StorageError(f"No se puede leer el archivo: {e}") from e
    except ValueError as e:
        raise StorageError(f"El archivo no es un proyecto válido: {e}") from e
    if not isinstance(data, dict) or data.get("format") not in (FORMAT, *OLD_FORMATS):
        raise StorageError("El archivo no es un documento de Dictador.")
    if int(data.get("version", 0)) > VERSION:
        raise StorageError("El proyecto es de una versión más nueva de Dictador.")
    paragraphs = [Paragraph.from_dict(p) for p in data.get("paragraphs", []) if isinstance(p, dict)]
    return paragraphs, data.get("mode", "prosa"), list(data.get("glossary", []))


# -- exportar ---------------------------------------------------------------------------------------------------------


def export_text(paragraphs: list[Paragraph]) -> str:
    return "\n\n".join(p.text.strip() for p in paragraphs if p.text.strip()) + "\n"


def export_markdown(paragraphs: list[Paragraph]) -> str:
    blocks = []
    for p in paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if p.style == "title":
            blocks.append("# " + " ".join(text.split()))
        elif p.mode == "poesia":
            blocks.append("  \n".join(line.strip() for line in text.split("\n")))
        else:
            blocks.append(text)
    return "\n\n".join(blocks) + "\n"


_LIST_ITEM = re.compile(r"^\s*(?:[-*•]|\d+[.)])\s+(.*)$")


def export_docx(path: str | Path, paragraphs: list[Paragraph]) -> None:
    from docx import Document

    doc = Document()
    for p in paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if p.style == "title":
            doc.add_heading(" ".join(text.split()), level=1)
            continue
        lines = text.split("\n")
        if p.mode == "notas" and any(_LIST_ITEM.match(line) for line in lines):
            for line in lines:
                m = _LIST_ITEM.match(line)
                if m:
                    numbered = line.strip()[0].isdigit()
                    doc.add_paragraph(m.group(1), style="List Number" if numbered else "List Bullet")
                elif line.strip():
                    doc.add_paragraph(line.strip())
            continue
        par = doc.add_paragraph()
        for i, line in enumerate(lines):
            run = par.add_run(line)
            if i < len(lines) - 1:
                run.add_break()
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def export(path: str | Path, paragraphs: list[Paragraph]) -> None:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        export_docx(path, paragraphs)
    elif suffix == ".md":
        _atomic_write(path, export_markdown(paragraphs))
    elif suffix == ".txt":
        _atomic_write(path, export_text(paragraphs))
    else:
        raise StorageError(f"Formato de exportación no admitido: {suffix}")


# -- importar ---------------------------------------------------------------------------------------------------------


def import_text(path: str | Path, mode: str = "prosa") -> list[Paragraph]:
    """Cada bloque separado por una línea en blanco entra como un párrafo editado a mano."""
    try:
        content = Path(path).read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        content = Path(path).read_text(encoding="cp1252")
    content = content.replace("\r\n", "\n").replace("\r", "\n")
    is_md = Path(path).suffix.lower() == ".md"
    blocks = [b.strip("\n") for b in re.split(r"\n\s*\n", content) if b.strip()]
    paragraphs = []
    for block in blocks:
        style = "normal"
        if is_md and re.match(r"^#{1,6}\s", block) and "\n" not in block:
            block = re.sub(r"^#{1,6}\s+", "", block)
            style = "title"
        if is_md:
            block = re.sub(r" {2,}\n", "\n", block)
        paragraphs.append(Paragraph(raw=block, status="edited", mode=mode, user_edited=True, style=style))
    return paragraphs
